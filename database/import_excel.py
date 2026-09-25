#!/usr/bin/env python3
"""Importe le classeur maître BVMAC dans PostgreSQL sans modifier l'ETL.

Chaque version distincte du XLSX devient un snapshot immuable. Le classeur est
lu en streaming avec openpyxl afin de limiter la RAM. Aucune règle métier,
normalisation ou correction n'est appliquée ici : les valeurs sont copiées telles
qu'elles existent dans l'Excel généré par l'ETL.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl
import psycopg
from psycopg.types.json import Jsonb

DEFAULT_DSN = os.environ.get(
    "BVMAC_IMPORT_DSN",
    "dbname=bvmac user=bvmacimport host=/var/run/postgresql",
)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def normalize_header(value: Any, index: int) -> str:
    text = str(value).strip() if value is not None else ""
    return text or f"column_{index}"


def import_workbook(path: Path, dsn: str = DEFAULT_DSN) -> int:
    if not path.is_file():
        raise FileNotFoundError(path)

    file_hash = sha256_file(path)
    stat = path.stat()
    source_mtime = datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc)

    with psycopg.connect(dsn) as conn:
        existing = conn.execute(
            "SELECT import_id, status FROM market.import_batch WHERE source_sha256=%s",
            (file_hash,),
        ).fetchone()
        if existing and existing[1] == "completed":
            print(f"Snapshot déjà importé: import_id={existing[0]} sha256={file_hash}")
            return int(existing[0])
        if existing:
            conn.execute("DELETE FROM market.excel_row WHERE import_id=%s", (existing[0],))
            conn.execute(
                "UPDATE market.import_batch SET status='loading', error_message=NULL, row_count=0, imported_at=now(), completed_at=NULL WHERE import_id=%s",
                (existing[0],),
            )
            import_id = int(existing[0])
        else:
            import_id = int(conn.execute(
                """
                INSERT INTO market.import_batch
                    (source_path, source_sha256, source_size_bytes, source_mtime, status)
                VALUES (%s,%s,%s,%s,'loading')
                RETURNING import_id
                """,
                (str(path), file_hash, stat.st_size, source_mtime),
            ).fetchone()[0])
        conn.commit()

        workbook = None
        total_rows = 0
        sheet_names: list[str] = []
        try:
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
            sheet_names = list(workbook.sheetnames)
            for sheet_name in sheet_names:
                sheet = workbook[sheet_name]
                rows = sheet.iter_rows(values_only=True)
                try:
                    first = next(rows)
                except StopIteration:
                    continue
                headers = [normalize_header(v, i + 1) for i, v in enumerate(first)]
                if len(headers) == 1 and headers[0].lower() == "(vide)":
                    continue

                batch: list[tuple[int, str, int, Jsonb]] = []
                for excel_row_number, values in enumerate(rows, start=2):
                    if not any(v is not None and str(v).strip() != "" for v in values):
                        continue
                    data = {
                        headers[i]: json_value(values[i]) if i < len(values) else None
                        for i in range(len(headers))
                    }
                    batch.append((import_id, sheet_name, excel_row_number, Jsonb(data)))
                    if len(batch) >= 500:
                        with conn.cursor() as cur:
                            cur.executemany(
                                "INSERT INTO market.excel_row(import_id,sheet_name,row_number,data) VALUES (%s,%s,%s,%s)",
                                batch,
                            )
                        total_rows += len(batch)
                        batch.clear()
                if batch:
                    with conn.cursor() as cur:
                        cur.executemany(
                            "INSERT INTO market.excel_row(import_id,sheet_name,row_number,data) VALUES (%s,%s,%s,%s)",
                            batch,
                        )
                    total_rows += len(batch)
                conn.commit()

            conn.execute(
                """
                UPDATE market.import_batch
                SET status='completed', completed_at=now(), workbook_sheets=%s, row_count=%s
                WHERE import_id=%s
                """,
                (Jsonb(sheet_names), total_rows, import_id),
            )
            conn.commit()
            print(json.dumps({
                "status": "completed",
                "import_id": import_id,
                "sha256": file_hash,
                "rows": total_rows,
                "sheets": sheet_names,
            }, ensure_ascii=False))
            return import_id
        except Exception as exc:
            conn.rollback()
            conn.execute(
                "UPDATE market.import_batch SET status='failed', error_message=%s WHERE import_id=%s",
                (str(exc)[:2000], import_id),
            )
            conn.commit()
            raise
        finally:
            if workbook is not None:
                workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", default="/var/lib/bvmac/output/bvmac_master.xlsx")
    parser.add_argument("--dsn", default=DEFAULT_DSN)
    args = parser.parse_args()
    try:
        import_workbook(Path(args.file), args.dsn)
    except Exception as exc:
        print(f"Import PostgreSQL échoué: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
