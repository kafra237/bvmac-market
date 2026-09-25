import importlib.util
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend" / "api"))
DSN = os.environ.get("BVMAC_TEST_POSTGRES_DSN")


def load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader
    spec.loader.exec_module(module)
    return module


@unittest.skipUnless(DSN and importlib.util.find_spec("psycopg") and importlib.util.find_spec("openpyxl"), "PostgreSQL d'intégration non configuré")
class PostgresIntegrationTests(unittest.TestCase):
    def setUp(self):
        import psycopg
        with psycopg.connect(DSN) as conn:
            conn.execute("TRUNCATE market.excel_row, market.import_batch RESTART IDENTITY CASCADE")
            conn.commit()

    def make_book(self, path: Path, close_price: float, result_net: float):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "dim_company"
        ws.append(["company_id", "isin", "ticker", "short_name", "full_name", "country", "sector", "listing_date", "currency"])
        ws.append([1, "CM0000000001", "TEST", "Test", "Test SA", "CM", "Finance", "2020-01-01", "XAF"])

        ws = wb.create_sheet("dim_opcvm")
        ws.append(["fund_id", "fund_name", "manager", "custodian", "category", "valuation_frequency", "initial_value", "inception_date_id"])

        ws = wb.create_sheet("fact_prices")
        ws.append(["company_id", "bulletin_date_id", "session_date_id", "status", "prev_price", "open_price", "close_price", "upper_limit", "lower_limit", "vol_bid", "vol_ask", "vol_traded", "value_traded", "num_transactions", "variation_pct", "daily_return_pct", "next_ref_price", "ytd_high", "ytd_low", "yoy_variation_pct", "last_div_amount", "last_div_date_id"])
        ws.append([1, 20260918, 20260918, "COTE", 12000, 12100, close_price, 13000, 11000, 1, 2, 3, 40000, 1, 1.5, 1.5, close_price, close_price, 10000, 2.1, 500, 20260501])

        ws = wb.create_sheet("fact_market_cap")
        ws.append(["company_id", "date_id", "close_price", "float_shares", "total_shares", "float_market_cap", "total_market_cap", "last_div_amount", "last_div_year", "last_div_date_id", "liquidity_pct", "eps", "per"])
        ws.append([1, 20260918, close_price, 1000, 2000, close_price * 1000, close_price * 2000, 500, 2026, 20260501, 5.5, 1000, 12.34])

        ws = wb.create_sheet("fact_index")
        ws.append(["date_id", "index_name", "index_value", "variation_day_pct"])
        ws.append([20260918, "BVMAC-AS", 100.25, 0.4])

        ws = wb.create_sheet("fact_opcvm_nav")
        ws.append(["fund_id", "bulletin_date_id", "nav", "nav_date_id", "prev_nav", "prev_nav_date_id", "var_prev_pct", "var_inception_pct"])

        ws = wb.create_sheet("fact_financials")
        ws.append(["company_id", "fiscal_year", "total_bilan", "capitaux_propres", "chiffre_affaires", "valeur_ajoutee", "resultat_net", "dividende_unitaire", "taux_rendement_brut_pct", "roe_publie_pct"])
        ws.append([1, 2025, 1000000, 400000, 700000, 300000, result_net, 550, 4.2, 11.8])
        wb.save(path)

    def test_snapshot_history_and_api_preserve_excel_values(self):
        importer = load_module("import_excel_test", ROOT / "database/import_excel.py")
        api = load_module("bvmac_api_test", ROOT / "backend/api/bvmac_api.py")
        api.DB_DSN = DSN
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder) / "bvmac_master.xlsx"
            self.make_book(p, 12345.67, 98765.43)
            first_id = importer.import_workbook(p, DSN)
            same_id = importer.import_workbook(p, DSN)
            self.assertEqual(first_id, same_id)

            self.make_book(p, 22345.89, 88765.21)
            second_id = importer.import_workbook(p, DSN)
            self.assertNotEqual(first_id, second_id)

            import psycopg
            with psycopg.connect(DSN) as conn:
                count = conn.execute("SELECT count(*) FROM market.import_batch WHERE status='completed'").fetchone()[0]
                self.assertEqual(count, 2)
                old_close = conn.execute(
                    "SELECT (data->>'close_price')::double precision FROM market.excel_row WHERE import_id=%s AND sheet_name='fact_prices'",
                    (first_id,),
                ).fetchone()[0]
                self.assertAlmostEqual(old_close, 12345.67)
                with self.assertRaises(psycopg.errors.RaiseException):
                    conn.execute("DELETE FROM market.excel_row WHERE import_id=%s", (first_id,))
                conn.rollback()

            body, _, _, _ = api._build_payload()
            payload = json.loads(body)
            self.assertEqual(payload["meta"]["storage"], "postgresql_from_excel")
            self.assertAlmostEqual(payload["prix"][0]["close_price"], 22345.89)
            self.assertAlmostEqual(payload["fin"][0]["resultat_net"], 88765.21)


if __name__ == "__main__":
    unittest.main()
